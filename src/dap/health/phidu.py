"""Read the layout of PHIDU Social Health Atlas sheets.

Each data sheet has block titles in row 1, optional sub-headings in row 2, time periods in row 4 and
measure labels in row 5. Row 2 is sometimes a sub-heading inside a block (Housing stress: mortgage,
rental) and sometimes the real title (the Census condition sheets, where row 1 is a caveat), so a
feature's `block` may match either, as long as exactly one column matches.
Area rows start at row 6: the PHAs first, then an "AUSTRALIA+" row, then state, SA4 and SA3 totals.
Thirty PHA codes are also SA3 codes (both have five digits), so rows must be split on the
AUSTRALIA+ marker, never on code length. Each state also has a pseudo-area coded <state>9999
("ABS cell adjustment" in Census sheets, "Unknown <state>" in admissions sheets); not a PHA.
"""

from dataclasses import dataclass
from pathlib import Path

import openpyxl

from dap.health.features import PHIDU_SOURCE, Feature

TITLE_ROW, SUBTITLE_ROW, PERIOD_ROW, MEASURE_ROW = 1, 2, 4, 5
TOTALS_MARKER = "AUSTRALIA+"
PSEUDO_AREA_SUFFIX = "9999"


def norm(value: object) -> str:
    return " ".join(str(value or "").split())


@dataclass(frozen=True)
class Column:
    index: int  # zero-based
    block: str  # row 1 title, carried right until the next one
    subblock: str  # row 2 sub-heading, carried right within the block
    period: str
    measure: str


def read_header(ws) -> list[Column]:
    rows = list(ws.iter_rows(min_row=1, max_row=MEASURE_ROW, values_only=True))
    title, subtitle = rows[TITLE_ROW - 1], rows[SUBTITLE_ROW - 1]
    period, measure = rows[PERIOD_ROW - 1], rows[MEASURE_ROW - 1]
    cols, block, subblock, block_period = [], "", "", ""
    for j in range(2, len(measure)):
        if j < len(title) and title[j]:
            block, subblock, block_period = norm(title[j]), "", ""
        if j < len(subtitle) and subtitle[j]:
            subblock, block_period = norm(subtitle[j]), ""
        if j < len(period) and period[j]:
            block_period = norm(period[j])
        if measure[j]:
            cols.append(Column(j, block, subblock, block_period, norm(measure[j])))
    return cols


@dataclass(frozen=True)
class AreaRows:
    """Area codes and raw cell values for one sheet, split at the AUSTRALIA+ marker."""

    pha: dict[str, tuple]  # PHA code -> row values, in sheet order
    pseudo: dict[str, tuple]  # <state>9999 adjustment or unknown-residence rows, never analysed
    totals: dict[str, tuple]  # codes after the marker (state, SA4 and SA3 totals); first row wins


def read_rows(ws) -> AreaRows:
    pha: dict[str, tuple] = {}
    pseudo: dict[str, tuple] = {}
    totals: dict[str, tuple] = {}
    after = False
    for row in ws.iter_rows(min_row=MEASURE_ROW + 1, values_only=True):
        if len(row) > 1 and norm(row[1]) == TOTALS_MARKER:
            after = True
            continue
        code = norm(row[0]) if row else ""
        if not code.isdigit():
            continue
        if after:
            totals.setdefault(code, row)
        elif len(code) == 5 and code.endswith(PSEUDO_AREA_SUFFIX):
            pseudo.setdefault(code, row)
        else:
            pha.setdefault(code, row)
    if not after:
        raise ValueError(f"sheet {ws.title!r} has no {TOTALS_MARKER} row")
    return AreaRows(pha, pseudo, totals)


def to_number(value: object) -> float | None:
    """PHIDU marks unavailable cells with #, .., n.p., n.a. and similar; all of them become None."""
    if isinstance(value, int | float):
        return float(value)
    try:
        return float(norm(value).replace(",", ""))
    except ValueError:
        return None


def locate(feature: Feature, header: list[Column]) -> Column:
    """Find the single column for a PHIDU feature, or raise KeyError."""
    if feature.source != PHIDU_SOURCE:
        raise ValueError(f"{feature.name} is not a PHIDU feature")
    hits = [
        c for c in header if feature.block in (c.block, c.subblock) and c.measure == feature.measure
    ]
    if len(hits) != 1:
        raise KeyError(
            f"{feature.name}: {len(hits)} columns match block {feature.block!r} and measure "
            f"{feature.measure!r} in sheet {feature.sheet!r}"
        )
    return hits[0]


def resolve_features(workbook: Path, features: tuple[Feature, ...]) -> dict[str, Column]:
    """Map each PHIDU feature to its column. Raises KeyError naming every missing feature."""
    wb = openpyxl.load_workbook(workbook, read_only=True)
    try:
        headers: dict[str, list[Column]] = {}
        found, problems = {}, []
        for f in features:
            if f.source != PHIDU_SOURCE:
                continue
            if f.sheet not in wb.sheetnames:
                problems.append(f"{f.name}: no sheet {f.sheet!r}")
                continue
            if f.sheet not in headers:
                headers[f.sheet] = read_header(wb[f.sheet])
            try:
                found[f.name] = locate(f, headers[f.sheet])
            except KeyError as e:
                problems.append(str(e).strip("'\""))
        if problems:
            raise KeyError("; ".join(problems))
        return found
    finally:
        wb.close()
